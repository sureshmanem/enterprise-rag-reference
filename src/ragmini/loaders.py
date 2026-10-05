"""Document loaders: PDF, Word, Markdown/text, and SharePoint (Microsoft Graph).

Best practices applied:
- PDFs are split per page so citations can point to a page number.
- Text is normalised (whitespace, hyphenated line breaks) before chunking.
- Empty / scanned-image pages are skipped and reported instead of silently indexed.
- SharePoint uses app-only (client credentials) auth; secrets come from env, never code.
"""

from __future__ import annotations

import logging
import os
import re
import tempfile
from collections.abc import Iterable, Iterator
from pathlib import Path

import httpx

from .models import Document

log = logging.getLogger(__name__)

SUPPORTED = {".pdf", ".docx", ".md", ".txt"}


def clean_text(text: str) -> str:
    text = text.replace("\x00", "")
    text = re.sub(r"(\w)-\n(\w)", r"\1\2", text)  # join hyphenated line breaks
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def load_pdf(path: Path) -> Iterator[Document]:
    from pypdf import PdfReader

    reader = PdfReader(str(path))
    for i, page in enumerate(reader.pages, start=1):
        text = clean_text(page.extract_text() or "")
        if not text:
            log.warning("No extractable text on %s page %d (scanned? consider OCR)", path, i)
            continue
        yield Document(text=text, source=path.name, metadata={"page": i, "type": "pdf"})


def load_docx(path: Path) -> Iterator[Document]:
    import docx

    d = docx.Document(str(path))
    parts: list[str] = []
    for para in d.paragraphs:
        if para.text.strip():
            # Keep headings visible so the chunker can split on them
            prefix = "## " if para.style.name.lower().startswith("heading") else ""
            parts.append(prefix + para.text)
    for table in d.tables:
        for row in table.rows:
            parts.append(" | ".join(c.text.strip() for c in row.cells))
    text = clean_text("\n\n".join(parts))
    if text:
        yield Document(text=text, source=path.name, metadata={"type": "docx"})


def load_text(path: Path) -> Iterator[Document]:
    text = clean_text(path.read_text(encoding="utf-8", errors="ignore"))
    if text:
        yield Document(text=text, source=path.name, metadata={"type": path.suffix.lstrip(".")})


def load_file(path: Path) -> Iterator[Document]:
    suffix = path.suffix.lower()
    if suffix == ".pdf":
        yield from load_pdf(path)
    elif suffix == ".docx":
        yield from load_docx(path)
    elif suffix in {".md", ".txt"}:
        yield from load_text(path)
    else:
        log.info("Skipping unsupported file %s", path)


def load_directory(root: str | Path) -> Iterator[Document]:
    root = Path(root)
    paths: Iterable[Path] = [root] if root.is_file() else sorted(root.rglob("*"))
    for p in paths:
        if p.is_file() and p.suffix.lower() in SUPPORTED:
            try:
                yield from load_file(p)
            except Exception:  # one bad file must not kill the whole ingestion run
                log.exception("Failed to load %s", p)


# ---------------------------------------------------------------- SharePoint


class SharePointLoader:
    """Downloads supported files from a SharePoint document library via Microsoft Graph.

    Env vars: SP_TENANT_ID, SP_CLIENT_ID, SP_CLIENT_SECRET (Entra ID app with
    Sites.Read.All application permission), plus a site hostname and path.
    """

    GRAPH = "https://graph.microsoft.com/v1.0"

    def __init__(self, hostname: str, site_path: str, folder: str = "") -> None:
        self.hostname, self.site_path, self.folder = hostname, site_path.strip("/"), folder.strip("/")
        self.tenant = os.environ["SP_TENANT_ID"]
        self.client_id = os.environ["SP_CLIENT_ID"]
        self.secret = os.environ["SP_CLIENT_SECRET"]

    def _token(self, client: httpx.Client) -> str:
        r = client.post(
            f"https://login.microsoftonline.com/{self.tenant}/oauth2/v2.0/token",
            data={
                "client_id": self.client_id,
                "client_secret": self.secret,
                "scope": "https://graph.microsoft.com/.default",
                "grant_type": "client_credentials",
            },
        )
        r.raise_for_status()
        return r.json()["access_token"]

    def load(self) -> Iterator[Document]:
        with httpx.Client(timeout=60) as client:
            client.headers["Authorization"] = f"Bearer {self._token(client)}"
            site = client.get(f"{self.GRAPH}/sites/{self.hostname}:/{self.site_path}").json()
            base = f"{self.GRAPH}/sites/{site['id']}/drive/root"
            url = f"{base}:/{self.folder}:/children" if self.folder else f"{base}/children"
            with tempfile.TemporaryDirectory() as tmp:
                while url:
                    page = client.get(url).json()
                    for item in page.get("value", []):
                        name = item.get("name", "")
                        if "file" not in item or Path(name).suffix.lower() not in SUPPORTED:
                            continue
                        dest = Path(tmp) / name
                        dest.write_bytes(
                            client.get(item["@microsoft.graph.downloadUrl"]).content
                        )
                        for doc in load_file(dest):
                            doc.metadata.update(
                                {"origin": "sharepoint", "web_url": item.get("webUrl")}
                            )
                            yield doc
                    url = page.get("@odata.nextLink")
